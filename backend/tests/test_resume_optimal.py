"""Optimal-resume tests: 1 page max, same format, JD-tailored order,
resume facts only, no upskill section."""
import io

import docx
import fitz

from app.services import latex_resume as lx
from app.services import resume_export as rx

LONG_RESUME = (
    "Jane Doe\njane@example.com | +91-98765 43210\n"
    "Frontend developer with 3 years experience building dashboards.\n"
    "## Education\nB.Sc Computer Science, Pune University, 2021\nM.Sc Data Science, 2023\n"
    "## Experience\n" + "".join(
        f"Frontend Engineer at Company{i}, built feature {i} with React and TypeScript\n"
        f"Shipped dashboard module {i}, Docker containers, AWS deploys\n" for i in range(10)) +
    "## Projects\n" + "".join(
        f"Project {i}: React app with CI/CD pipeline and Kubernetes manifests\n" for i in range(8)) +
    "## Certifications\nAWS Cloud Practitioner 2024\nScrum Master 2023\n"
    "## Achievements\nWon hackathon 2023\nDean's list 2021\n"
)

JD_REQ = ["Docker", "AWS", "Kubernetes"]


def _data():
    out = {
        "optimized_resume": {"summary": "Frontend dev skilled in React."},
        "_candidate": ["React", "JavaScript", "TypeScript", "Docker"],
        "_required": JD_REQ,
        "_resume_text": LONG_RESUME,
        "skill_gap": {"strong": ["React"], "partial": [],
                      "missing": [{"skill": "Go", "criticality": "High",
                                   "reason": "r", "weeks_to_learn": 3}],
                      "critical_missing": ["Go"]},
    }
    return rx.build_export_data(out)


def test_no_upskill_section_anywhere():
    data = _data()
    tex = lx.build_latex(data)
    assert "Upskill" not in tex and "upskill" not in tex.replace("areas_to_upskill", "")
    pdf = rx.to_pdf(data)
    text = "\n".join(p.get_text() for p in fitz.open(stream=pdf, filetype="pdf"))
    assert "Upskill" not in text
    docx = rx.to_docx(data)
    assert b"Upskill" not in docx


def test_one_page_guarantee_long_resume():
    data = _data()
    pdf, engine, _tex = lx.build_one_page_pdf(data)
    assert pdf is not None and engine in {"latex", "latex-tight"}
    assert pdf.startswith(b"%PDF")
    assert len(fitz.open(stream=pdf, filetype="pdf")) == 1


def test_jd_relevant_bullets_first():
    data = _data()
    tex = lx.build_latex(data)
    first_docker = tex.index("Docker")
    assert first_docker < tex.index("Experience") + 3000  # near the top
    # Experience section leads with JD-matching lines.
    exp = tex.split("Experience")[1].split("Skills Summary")[0]
    first_item = exp.split("resumeItemWithoutTitle")[1]
    assert "Docker" in first_item or "AWS" in first_item or "Kubernetes" in first_item


def test_only_resume_facts_rendered():
    data = _data()
    tex = lx.build_latex(data)
    assert "Go" not in tex  # missing skill never appears on the resume
    squashed = tex.replace(" ", "").replace("\\", "")
    # Kept facts (education uncapped, top JD-ranked lines).
    for token in ["PuneUniversity", "Project5", "Docker"]:
        assert token in squashed, token
    # Low-priority overflow honestly capped for 1 page — never invented.
    assert "Company3" not in tex


def test_prioritize_is_stable_and_capped():
    lines = ["plain line one", "Docker deploy pipeline", "plain line two",
             "AWS hosted setup"] + [f"filler {i}" for i in range(20)]
    out = lx.prioritize_lines(lines, ["Docker", "AWS"])
    assert out[0] == "Docker deploy pipeline" and out[1] == "AWS hosted setup"
    assert out[2:] == ["plain line one", "plain line two"] + [f"filler {i}" for i in range(20)]
    tex_lines = lx.build_latex({**_data(), "sections": {
        "education": [], "experience": lines, "projects": [],
        "certifications": [], "activities": []}})
    assert tex_lines.count("\\resumeItemWithoutTitle") <= 7


READABLE = (
    "Jane Doe\njane@example.com | +91-98765 43210\n"
    "## Education\nPune University\nPune, Maharashtra\nB.Tech CSE, 2024\n"
    "## Projects\nPortal - React dashboard\nGitHub\n"
    "\u2022 Built dashboard with React\n- Wrote 20+ SQL queries\n"
    "## Skills\nLanguages: Java, Python\nCloud: AWS, S3\n"
)


def _readable_data():
    return rx.build_export_data(
        {"_resume_text": READABLE, "_candidate": ["Java", "Python", "AWS", "S3"],
         "_required": ["Java"], "optimized_resume": {}, "skill_gap": {}},
        resume_text=READABLE)


def test_skills_render_one_group_per_line():
    data = _readable_data()
    assert [g["label"] for g in data["skill_groups"]] == ["Languages", "Cloud"]
    tex = lx.build_latex(data)
    assert "Languages:" in tex and "Cloud:" in tex
    skills_tex = tex.split("Skills Summary")[1]
    assert "Java, Python, AWS, S3" not in skills_tex  # never one comma soup line
    pdf_text = "\n".join(p.get_text() for p in fitz.open(stream=rx.to_pdf(data), filetype="pdf"))
    assert "Languages: Java, Python" in pdf_text
    document = docx.Document(io.BytesIO(rx.to_docx(data)))
    assert "Languages: Java, Python" in [p.text for p in document.paragraphs]
    assert "Cloud: AWS, S3" in [p.text for p in document.paragraphs]


PDF_STYLE = (
    "Aryan Akhare\nFull Stack Developer\naryan@example.com\n"
    "## Projects\nDSA Visualizer\nOct 2024\n"
    "\u2013 Built an interactive learning platform using React 18 and Vite, delivering "
    "15+ animated visualizations across 6 core data\n"
    "structures and algorithms.\n"
    "\u2013 Led a 5-member team by coordinating task allocation.\n"
    "Anandvan E-Commerce Platform\nMar 2025\n"
    "\u2013 Developed a full-stack e-commerce platform using React.\n"
    "## Awards and Certifications\nData Engineering Virtual Internship\n"
    "Jul 2025 \u2013 Sep 2025\nAICTE NEAT \u2013 EduSkills \u2013 AWS Academy\n"
    "## Education\nPune University\n2024\nB.Tech CSE\n"
)


def test_pdf_en_dash_bullets_and_dates_parse_as_entries():
    data = rx.build_export_data({"_resume_text": PDF_STYLE})
    projects = data["blocks"]["projects"]
    assert [p["title"] for p in projects][:2] == [
        "DSA Visualizer", "Anandvan E-Commerce Platform"]
    assert projects[0]["date"] == "Oct 2024"
    assert projects[0]["bullets"][0].startswith("Built an interactive")
    assert "algorithms." in projects[0]["bullets"][0]  # wrapped line joined
    assert projects[1]["bullets"] == [
        "Developed a full-stack e-commerce platform using React."]
    cert = data["blocks"]["certifications"][0]
    assert cert["title"] == "Data Engineering Virtual Internship"
    assert cert["date"].startswith("Jul 2025")
    assert cert["details"] == ["AICTE NEAT \u2013 EduSkills \u2013 AWS Academy"]
    tex = lx.build_latex(data)
    assert "\\resumeProjectHeading{DSA Visualizer}{Oct 2024}" in tex
    assert "Awards and Certifications" not in tex


def test_education_grouped_and_bullets_have_no_double_glyph():
    data = _readable_data()
    edu = data["sections"]["education"][0]
    assert edu["title"] == "Pune University"
    assert any("B.Tech CSE" in d for d in edu["details"])
    project = data["sections"]["projects"][0]
    assert project["title"].startswith("Portal")
    assert project["bullets"] == ["Built dashboard with React", "Wrote 20+ SQL queries"]
    document = docx.Document(io.BytesIO(rx.to_docx(data)))
    texts = [p.text for p in document.paragraphs]
    assert "Built dashboard with React" in texts
    assert not any(t.startswith(("\u2022", "- ")) for t in texts)


FULL_STACK = (
    "Aarsh Vaidya\n"
    "+91 77210 80610 | aarsh@example.com | linkedin.com/in/aarsh37 | github.com/Aarsh-37\n"
    "Summary\n"
    "Full-stack developer skilled in the MERN stack, with a strong foundation in machine learning.\n"
    "Education\n"
    "Shri Ramdeobaba College of Engineering and Management, Nagpur\n"
    "2023 \u2013 2027\n"
    "B.Tech in Artificial Intelligence & Machine Learning\n"
    "CGPA: 8.69\n"
    "Tip Top Convent, Nagpur\n"
    "2022\n"
    "Class XII (HSC)\n"
    "84.83%\n"
    "Saraswati Vidyalaya, Nagpur\n"
    "2020\n"
    "Class X (SSC)\n"
    "95.4%\n"
    "Experience\n"
    "IIT Roorkee Summer Training & Internship\n"
    "| Certificate\n"
    "Jun 2026 \u2013 Jul 2026\n"
    "E&ICT Academy, IIT Roorkee\n"
    "(Remote)\n"
    "\u2022 Completed hands-on training in Generative AI, prompt engineering, and multi-agent "
    "orchestration using Ollama (local LLMs) and\n"
    "CrewAI, culminating in a full-stack capstone project.\n"
    "Projects\n"
    "Kisan Mitra - Smart Crop Advisory System | GitHub\n"
    "Feb 2026 \u2013 Apr 2026\n"
    "Next.js, FastAPI, PostgreSQL\n"
    "\u2013 Architected a microservices platform with Dockerized deployment.\n"
    "Skills\n"
    "Languages: Java, Python\n"
    "Certifications\n"
    "Oracle Agentic AI Certified Foundations Associate \u2013 Oracle University\n"
    "Jul 2026\n"
    "Deep Learning Certification \u2013 NVIDIA\n"
    "Feb 2026\n"
    "Co-Curricular Activities\n"
    "\u2022 State-Level Cricket Player, Active in college competitions.\n"
)


def test_summary_heading_is_not_used_as_job_title():
    data = rx.build_export_data({"_resume_text": FULL_STACK})
    assert data["title"] != "Summary"
    assert data["summary"].startswith("Full-stack developer")
    assert len(data["sections"]["summary"]) >= 1
    tex = lx.build_latex(data)
    assert "\\textsc{SUMMARY}" not in tex


def test_full_stack_education_experience_and_certs_group_correctly():
    data = rx.build_export_data({"_resume_text": FULL_STACK})
    edu = data["blocks"]["education"]
    assert [e["title"] for e in edu] == [
        "Shri Ramdeobaba College of Engineering and Management, Nagpur",
        "Tip Top Convent, Nagpur",
        "Saraswati Vidyalaya, Nagpur",
    ]
    assert edu[0]["date"] == "2023 \u2013 2027"
    assert edu[1]["date"] == "2022" and edu[2]["date"] == "2020"
    assert edu[0]["details"] == [
        "B.Tech in Artificial Intelligence & Machine Learning", "CGPA: 8.69"]

    exp = data["blocks"]["experience"]
    assert len(exp) == 1
    assert exp[0]["title"] == "IIT Roorkee Summer Training & Internship"
    assert exp[0]["date"].startswith("Jun 2026")
    assert exp[0]["details"] == ["E&ICT Academy, IIT Roorkee", "(Remote)"]
    assert len(exp[0]["bullets"]) == 1  # wrapped PDF line rejoined
    assert exp[0]["bullets"][0].endswith("capstone project.")

    certs = data["blocks"]["certifications"]
    assert [c["title"] for c in certs] == [
        "Oracle Agentic AI Certified Foundations Associate \u2013 Oracle University",
        "Deep Learning Certification \u2013 NVIDIA",
    ]
    assert [c["date"] for c in certs] == ["Jul 2026", "Feb 2026"]
    activities = data["blocks"]["activities"]
    assert activities[0]["bullets"] == [
        "State-Level Cricket Player, Active in college competitions."]


def test_tech_stack_renders_as_italic_detail_line():
    data = rx.build_export_data({"_resume_text": FULL_STACK})
    tex = lx.build_latex(data)
    assert "\\resumeDetailLine{Next.js, FastAPI, PostgreSQL}" in tex
    assert "\\resumeDetailLine{E\\&ICT Academy, IIT Roorkee | (Remote)}" in tex
    pdf_text = "\n".join(
        p.get_text() for p in fitz.open(stream=rx.to_pdf(data), filetype="pdf"))
    for token in ("Tip Top Convent", "84.83%", "Deep Learning Certification",
                  "Next.js, FastAPI, PostgreSQL"):
        assert token in pdf_text
    document = docx.Document(io.BytesIO(rx.to_docx(data)))
    texts = [p.text for p in document.paragraphs]
    assert "Next.js, FastAPI, PostgreSQL" in texts
    assert "State-Level Cricket Player, Active in college competitions." in texts
